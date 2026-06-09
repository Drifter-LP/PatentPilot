"""Generate PatentPilot project documentation as Word (.docx)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUTPUT = Path(__file__).resolve().parent / "PatentPilot项目说明.docx"


def set_cn_font(run, name: str = "微软雅黑", size_pt: int | None = None, bold: bool = False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    run.bold = bold


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_cn_font(run, size_pt=22, bold=True)
    run.font.color.rgb = RGBColor(0x0A, 0x0A, 0x0A)


def add_subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_cn_font(run, size_pt=12)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(level=level)
    h.clear()
    run = h.add_run(text)
    sizes = {1: 16, 2: 14, 3: 12}
    set_cn_font(run, size_pt=sizes.get(level, 12), bold=True)
    run.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)


def add_para(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_cn_font(r1, size_pt=11, bold=True)
    r2 = p.add_run(text)
    set_cn_font(r2, size_pt=11)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(6)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_cn_font(run, size_pt=11)
    p.paragraph_format.line_spacing = 1.35


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                set_cn_font(run, size_pt=10, bold=True)
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, cell_text in enumerate(row):
            row_cells[c_idx].text = cell_text
            for p in row_cells[c_idx].paragraphs:
                for run in p.runs:
                    set_cn_font(run, size_pt=10)
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    add_title(doc, "PatentPilot 项目说明文档")
    add_subtitle(doc, "专利潜力初筛与导航工具 · 产品说明与技术概要")
    add_subtitle(doc, "版本：Demo v1.0  |  更新日期：2026 年 6 月")
    doc.add_paragraph()

    add_heading(doc, "一、产品定位")
    add_para(
        doc,
        "PatentPilot 是一款面向专利申请前置阶段的 AI 辅助工具。在用户尚未投入大量专利检索成本、"
        "尚未形成完整技术方案或权利要求书之前，帮助用户用自然语言快速评估「想法或科研成果是否值得"
        "继续往专利方向推进」，并给出可理解的改进方向与风险提示。",
    )
    add_para(doc, "工具输出为解释型分析结论，用于导航与决策参考，不构成正式法律意见。", bold_prefix="定位边界：")
    add_para(doc, "", bold_prefix="目标用户：")
    add_bullet(doc, "创新小白 / 创业者：有初步 idea，缺乏专利规划经验，需要低门槛的可行性判断。")
    add_bullet(doc, "科研工作者 / 高校团队：已有论文摘要、项目描述或阶段性成果，需要将科研成果评估为专利化路径。")
    add_para(doc, "", bold_prefix="产品 Slogan：")
    add_para(doc, "检验你的想法是不是一个申请专利的金点子！")

    add_heading(doc, "二、产品创新与差异点")
    add_heading(doc, "2.1 与智慧芽等平台的差异", level=2)
    add_para(
        doc,
        "智慧芽等专利检索平台主要解决的是：在用户已有较成熟技术方案、并愿意投入检索资源之后，"
        "查询相关领域是否已存在相同或近似的专利技术——即「有没有撞车」的问题。",
    )
    add_para(
        doc,
        "PatentPilot 则聚焦更早的阶段：在正式申请与深度检索之前，验证想法或成果的可行性、"
        "识别主要风险与改进方向——即「值不值得做、怎么做更好」的问题。",
    )
    add_table(
        doc,
        ["对比维度", "智慧芽等检索平台", "PatentPilot"],
        [
            ["使用阶段", "成熟方案后的重复性检索", "申请前的前置可行性初筛"],
            ["核心问题", "是否已有相同/近似专利", "是否值得专利化、如何改进"],
            ["交互方式", "检索式、分类号等专业操作", "自然语言描述或粘贴摘要"],
            ["用户门槛", "较高，面向专业人士", "较低，面向个人创新者"],
            ["商业倾向", "以 2B 企业/机构为主", "更面向个人与小团队轻量验证"],
            ["输出形态", "专利列表、法律状态、引证关系", "结构化报告、评分、行动建议"],
        ],
    )
    add_para(doc, "两者定位互补：PatentPilot 不替代正式检索，而是帮助用户在更早阶段做出是否继续投入的决策。")

    add_heading(doc, "2.2 双模式创新设计", level=2)
    add_table(
        doc,
        ["模式", "面向人群", "输入方式", "输出形态"],
        [
            ["💡 Idea Mode", "非专业创新者", "自然语言描述创新想法", "Bento 仪表盘 + 五维雷达图 + 行动建议"],
            ["🔬 Research Mode", "专业科研人员", "论文摘要 / 项目描述", "科研风报告 § I–V + PPI 多维评估"],
        ],
    )

    add_heading(doc, "三、项目内容")
    add_heading(doc, "3.1 功能模块", level=2)
    add_bullet(doc, "开场页：品牌展示、Slogan、Get Started 入口、产品说明；单屏无滚动，一次点击进入。")
    add_bullet(doc, "输入页：Idea / Research 双模式切换、示例卡片一键填入、侧边栏新建分析与模式标识。")
    add_bullet(doc, "分析引擎：基于 DeepSeek 大模型，按模式调用不同 Prompt，返回结构化 JSON。")
    add_bullet(doc, "进度体验：全屏进度遮罩、分阶段文案与进度条，后台线程执行分析。")
    add_bullet(doc, "Idea 结果页：专利潜力评分、雷达图、机会/风险/建议分区、技术领域标签。")
    add_bullet(doc, "Research 结果页：PPI 指数、§ I 综合结论、§ II 创新性、§ III 风险、§ IV 多维评估、§ V 策略建议。")

    add_heading(doc, "3.2 用户流程", level=2)
    add_para(doc, "开场页 → 选择模式并输入内容 → AI 分析（进度遮罩）→ 查看结构化报告 → 新建分析或返回修改输入。")

    add_heading(doc, "四、系统设计")
    add_heading(doc, "4.1 技术栈", level=2)
    add_table(
        doc,
        ["层级", "技术选型"],
        [
            ["前端 / UI", "Streamlit 1.32+"],
            ["后端语言", "Python 3.10+"],
            ["大模型", "DeepSeek（ModelArts MaaS 兼容接口，可配置）"],
            ["配置管理", "python-dotenv + .env；Streamlit Cloud Secrets"],
            ["数值计算", "NumPy（Idea 模式图表，延迟加载）"],
        ],
    )

    add_heading(doc, "4.2 系统架构", level=2)
    add_para(
        doc,
        "采用 Streamlit 单页应用架构。app.py 负责路由与页面编排；PatentAnalyzer 封装 LLM 调用与 JSON 解析；"
        "templates.py 定义 Idea / Research 双 Prompt 契约；dashboard.py 与 research_dashboard.py 分别渲染"
        "两类结果报告。开场页与主应用分离，并通过延迟 import 优化冷启动性能。",
    )
    add_para(doc, "", bold_prefix="核心模块：")
    add_bullet(doc, "app.py — 主入口、开场页、输入页、结果路由")
    add_bullet(doc, "src/services/analyzer.py — 分析编排")
    add_bullet(doc, "src/prompts/templates.py — Prompt 与 JSON Schema")
    add_bullet(doc, "src/api/client.py — DeepSeek HTTP 客户端")
    add_bullet(doc, "src/ui/analysis_progress.py — 进度遮罩")
    add_bullet(doc, "src/ui/dashboard.py — Idea Mode 仪表盘")
    add_bullet(doc, "src/ui/research_dashboard.py — Research Mode 报告")

    add_heading(doc, "4.3 安全与部署", level=2)
    add_bullet(doc, "API Key 通过环境变量或 Streamlit Secrets 注入，不在界面暴露输入框。")
    add_bullet(doc, "支持本地运行与 Streamlit Community Cloud 云端部署。")

    add_heading(doc, "五、当前完成情况")
    add_heading(doc, "5.1 已完成功能", level=2)
    done_items = [
        "Streamlit 完整主流程（开场 → 输入 → 分析 → 报告）",
        "Idea Mode / Research Mode 双模式与 DeepSeek 接入",
        "Idea 仪表盘 v3（Bento + 雷达图 + 结论与建议）",
        "Research 科研风报告（§ I–V、PPI、五维评估、sticky 顶栏）",
        "分析进度遮罩与全视口浅色蒙层",
        "开场页布局优化、单击进入、延迟 import 性能优化",
        "Secrets / 环境变量配置，适配云端部署",
        "设计稿与 preview 预览资源",
    ]
    for item in done_items:
        add_bullet(doc, item)

    add_heading(doc, "5.2 已知限制", level=2)
    add_bullet(doc, "热力图及部分竞争数据为 Demo 示意，非真实专利库检索结果。")
    add_bullet(doc, "历史记录仅在当前会话内展示，尚未持久化存储。")
    add_bullet(doc, "分析基于大模型通用知识，不能替代官方检索与律师意见。")

    add_heading(doc, "六、运行效果")
    add_para(
        doc,
        "部署并配置 DEEPSEEK_API_KEY 后，用户可在浏览器中完成完整体验："
        "开场页展示品牌与 Slogan；进入主应用后选择 Idea 或 Research 模式输入内容；"
        "分析过程中显示全屏进度遮罩；Idea 模式返回直观 Bento 仪表盘，Research 模式返回"
        "分节式学术报告，包含 PPI 评分、创新性/风险分析、多维评估矩阵与策略建议列表。",
    )
    add_para(doc, "", bold_prefix="本地启动：")
    add_para(doc, "streamlit run app.py  →  默认 http://localhost:8501")
    add_para(doc, "", bold_prefix="Research 报告预览：")
    add_para(doc, "streamlit run preview/research_dashboard_demo.py --server.port 8502")

    add_heading(doc, "七、后续计划")
    add_heading(doc, "7.1 功能扩展", level=2)
    add_bullet(doc, "接入真实专利检索 API，替换示意性 prior art 与竞争数据。")
    add_bullet(doc, "分析结果导出（PDF / Markdown / Word）。")
    add_bullet(doc, "历史记录持久化（SQLite 或本地文件）。")
    add_bullet(doc, "多轮对话式 Agent 澄清，追问技术细节后再生成报告。")

    add_heading(doc, "7.2 体验与工程", level=2)
    add_bullet(doc, "精简开场 iframe，进一步缩短冷启动时间。")
    add_bullet(doc, "单元测试与 CI（mock LLM smoke test）。")
    add_bullet(doc, "Docker 化与一键部署文档。")
    add_bullet(doc, "演示 GIF 与英文版 README。")

    add_heading(doc, "八、免责声明")
    add_para(
        doc,
        "PatentPilot Demo 仅供产品演示、教学与技术验证。分析结果基于大模型通用知识生成，"
        "不能替代专利律师或官方检索系统的专业意见。正式申请专利前，请咨询具备资质的专业机构"
        "并完成官方检索与自由实施（FTO）评估。",
    )

    return doc


def main() -> None:
    doc = build_document()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    main()
